from pathlib import Path
from datetime import date
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Dopilot_Propuesta_Integral.docx"

BLACK = "171717"
MUTED = "5F6368"
LIGHT = "F5F6F7"
MID = "D9DDE1"
TEAL = "0F766E"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=MID, size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_keep_with_next(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    properties.append(keep)


def remove_paragraph_borders(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def set_run_color(run, color=BLACK):
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        if widths:
            cell.width = Cm(widths[index])
        set_cell_shading(cell, BLACK)
        set_cell_borders(cell)
        set_cell_margins(cell, 120, 140, 120, 140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        run.bold = True
        run.font.size = Pt(9)
        set_run_color(run, WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            if widths:
                cell.width = Cm(widths[index])
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9)
            set_run_color(run, BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    set_run_color(run)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        lead.font.size = Pt(10.5)
        set_run_color(lead, BLACK)
        rest = text[len(bold_lead):]
        run = paragraph.add_run(rest)
        run.font.size = Pt(10.5)
        set_run_color(run, MUTED)
    else:
        run = paragraph.add_run(text)
        run.font.size = Pt(10.5)
        set_run_color(run, MUTED)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(6)
    set_keep_with_next(paragraph)
    for run in paragraph.runs:
        set_run_color(run, BLACK)
    return paragraph


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(MUTED)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(8)
    title_paragraph_properties = title.element.get_or_add_pPr()
    title_borders = title_paragraph_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_paragraph_properties.remove(title_borders)

    for name, size in (("Heading 1", 17), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)

    for name in ("List Bullet", "List Bullet 2"):
        style = styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(MUTED)


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("DOPILOT")
    run.bold = True
    run.font.size = Pt(9)
    set_run_color(run, TEAL)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("Dopilot | Propuesta integral de producto y negocio  ·  ")
    run.font.size = Pt(8)
    set_run_color(run, MUTED)
    add_field(footer, "PAGE")

    title = doc.add_paragraph(style="Title")
    title.add_run("Dopilot Propuesta Integral de Producto y Negocio")
    remove_paragraph_borders(title)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Modelo SaaS para automatizar la operación de consultorios médicos")
    run.font.size = Pt(13)
    run.italic = True
    set_run_color(run, TEAL)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(18)
    run = meta.add_run(f"Documento de definición comercial y operativa  |  {date.today().strftime('%d de %B de %Y')}")
    run.font.size = Pt(9.5)
    set_run_color(run, MUTED)

    add_heading(doc, "Resumen ejecutivo", 1)
    add_body(doc, "Dopilot es un SaaS para doctores y consultorios que centraliza agenda, pacientes, servicios y comunicación automatizada. Su diferenciador no es únicamente permitir reservar una cita: es reducir el trabajo repetitivo de recepción y mantener al paciente acompañado antes y después de la consulta.")
    add_body(doc, "El modelo comercial base es simple: cada doctor paga una licencia de $800 MXN al mes. La licencia incluye $150 MXN de crédito mensual para mensajes y automatizaciones operadas mediante Sent.dm. Cuando el consumo supera ese crédito, el excedente se factura como cargo variable adicional, con desglose visible para el doctor.")
    add_body(doc, "La primera etapa debe concentrarse en una experiencia confiable de agenda y automatizaciones de mensajería. Los agentes de voz con Vapi quedan como una expansión posterior, detrás de una frontera de integración segura y después de validar demanda, costos y consentimiento del paciente.")

    add_heading(doc, "Producto y propuesta de valor", 1)
    add_heading(doc, "Qué resuelve Dopilot", 2)
    for item in [
        "Permite que un doctor publique sus servicios, horarios y disponibilidad sin depender de hojas de cálculo o múltiples canales.",
        "Ayuda al paciente a encontrar al profesional adecuado, reservar y consultar el estado de sus citas.",
        "Reduce tareas manuales mediante recordatorios, confirmaciones y notificaciones automatizadas.",
        "Mantiene la información clínica y operativa separada por usuario y rol, preparando una evolución hacia consultorios multiusuario.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Por qué un doctor pagaría", 2)
    add_body(doc, "El doctor no compra solamente una agenda. Compra tiempo operativo recuperado, menos conversaciones repetitivas y una experiencia más consistente para sus pacientes. La comunicación automatizada debe presentarse como una herramienta de continuidad de atención, no como spam ni como un sustituto del criterio clínico.")

    add_heading(doc, "Alcance del producto", 1)
    add_table(
        doc,
        ["Área", "Capacidad inicial", "Resultado para el doctor"],
        [
            ["Agenda", "Horarios, disponibilidad, reservas, cancelaciones y reprogramaciones", "Menos coordinación manual"],
            ["Pacientes", "Listado, ficha básica e historial disponible según permisos", "Contexto antes de atender"],
            ["Servicios", "Nombre, duración, precio y estado activo", "Oferta clara y actualizable"],
            ["Automatizaciones", "Recordatorios y notificaciones con Sent.dm", "Seguimiento constante"],
            ["Experiencia", "Panel web con navegación compacta y estados claros", "Operación rápida y profesional"],
        ],
        widths=[3.0, 9.3, 5.5],
    )
    add_body(doc, "Fuera del alcance inicial quedan el diagnóstico automático, la prescripción, la toma de decisiones clínicas y cualquier promesa de atención médica autónoma. Dopilot automatiza coordinación y comunicación; el doctor conserva la responsabilidad clínica.")

    add_heading(doc, "Modelo de negocio", 1)
    add_heading(doc, "Licencia por doctor", 2)
    add_table(
        doc,
        ["Concepto", "Regla comercial", "Nota operativa"],
        [
            ["Licencia Dopilot", "$800 MXN por doctor al mes", "Cobro recurrente por cuenta activa"],
            ["Crédito Sent.dm incluido", "$150 MXN por doctor al mes", "Se consume únicamente en mensajes y automatizaciones elegibles"],
            ["Excedente Sent.dm", "Cargo variable adicional", "Se factura el consumo que supere el crédito incluido"],
            ["Usuarios adicionales", "Pendiente de definir", "No incluirlos en el precio base hasta validar el modelo multiusuario"],
            ["Agentes de voz Vapi", "No incluidos en la licencia inicial", "Producto futuro con precio y consumo propios"],
        ],
        widths=[4.2, 5.3, 8.3],
    )
    add_heading(doc, "Reglas de facturación recomendadas", 2)
    for item in [
        "El crédito de $150 MXN es de uso operativo: no se convierte en saldo en efectivo, no es transferible y vence al terminar el ciclo mensual.",
        "El consumo de Sent.dm debe registrarse por doctor, tipo de automatización y periodo de facturación.",
        "El doctor debe poder consultar consumo acumulado, crédito disponible y excedente estimado desde Dopilot.",
        "El cargo excedente se calcula con la tarifa vigente del proveedor y la política comercial publicada por Dopilot; cualquier ajuste de tarifa debe comunicarse antes del siguiente ciclo.",
        "Dopilot debe establecer límites de seguridad para evitar consumos accidentales o automatizaciones en bucle.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Ejemplo económico simple", 2)
    add_table(
        doc,
        ["Concepto", "Importe mensual"],
        [
            ["Ingreso por licencia", "$800 MXN"],
            ["Crédito de mensajería incluido", "Hasta $150 MXN en Sent.dm"],
            ["Ingreso disponible antes de otros costos", "$650 MXN"],
            ["Consumo que exceda $150 MXN", "Se suma como cargo variable"],
        ],
        widths=[11.5, 6.3],
    )
    add_body(doc, "Este ejemplo no representa margen neto. Antes de fijar una meta financiera deben descontarse impuestos, procesamiento de pagos, infraestructura, soporte, costos de desarrollo y cualquier tarifa comercial de Sent.dm u otros proveedores.")

    add_heading(doc, "Flujo operativo de Dopilot", 1)
    add_table(
        doc,
        ["Momento", "Acción del paciente", "Respuesta del sistema"],
        [
            ["Descubrimiento", "Busca un doctor o entra por un enlace", "Muestra perfil, servicios y horarios disponibles"],
            ["Reserva", "Selecciona servicio, fecha y hora", "Valida disponibilidad y crea la cita"],
            ["Confirmación", "Recibe y responde el mensaje", "Envía recordatorio y registra el estado"],
            ["Cambio", "Solicita cancelar o reprogramar", "Aplica reglas de agenda y notifica al doctor"],
            ["Atención", "Acude a la consulta", "Conserva el historial operativo de la cita"],
            ["Seguimiento", "Recibe una solicitud de valoración o próxima acción", "Automatiza el mensaje sin intervenir en decisiones clínicas"],
        ],
        widths=[3.2, 7.0, 7.6],
    )
    add_body(doc, "Sent.dm es el canal de automatización de mensajería en la etapa actual. La integración debe trabajar con plantillas aprobadas, consentimiento y límites de frecuencia. Vapi podrá añadirse después para llamadas de voz, pero no debe introducirse como dependencia central antes de tener contratos, autenticación y observabilidad definidos.")

    add_heading(doc, "Cliente ideal y posicionamiento", 1)
    add_table(
        doc,
        ["Segmento", "Necesidad principal", "Mensaje comercial"],
        [
            ["Doctor independiente", "Agenda y recordatorios sin contratar recepción adicional", "Tu consultorio sigue atendiendo aunque tú estés en consulta"],
            ["Consultorio pequeño", "Ordenar pacientes, horarios y mensajes", "Un solo lugar para operar el día"],
            ["Especialidades con citas de alto valor", "Reducir huecos y ausencias", "Cada cita confirmada protege tus ingresos"],
            ["Clínica multi-doctor", "Roles, permisos y varias agendas", "Siguiente expansión del producto"],
        ],
        widths=[4.2, 7.0, 6.6],
    )
    add_body(doc, "La entrada comercial más clara es el doctor que ya recibe consultas por WhatsApp, llamadas o redes sociales y pierde tiempo coordinando horarios. La promesa debe ser medible: más citas confirmadas, menos trabajo repetitivo y visibilidad del consumo de automatizaciones.")

    add_heading(doc, "Arquitectura y operación técnica", 1)
    for item in [
        "Dopilot: aplicación web para doctores y pacientes, con agenda, servicios, pacientes, notificaciones y autenticación por rol.",
        "Sent.dm: proveedor de mensajería y automatizaciones de la primera etapa; su consumo se registra para facturación por doctor.",
        "Dominio de citas: disponibilidad, duración, conflictos, cancelación y reprogramación deben mantenerse en servicios de dominio compartidos.",
        "Integraciones futuras: Vapi debe entrar como adaptador aislado, con autenticación, límites, trazabilidad y confirmación explícita antes de escribir en la agenda.",
        "Datos: el crecimiento hacia clínicas requiere tenant, membresías, permisos y zona horaria antes de vender cuentas multiusuario.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Roadmap recomendado", 1)
    add_table(
        doc,
        ["Fase", "Objetivo", "Criterio de salida"],
        [
            ["1. Producto base", "Consolidar agenda, perfiles, servicios, pacientes y experiencia visual", "Un doctor puede operar su día completo sin soporte manual"],
            ["2. Sent.dm", "Recordatorios, confirmaciones, plantillas y medidor de consumo", "Cada mensaje tiene trazabilidad y costo asignable"],
            ["3. Comercialización", "Pagos recurrentes, estado de licencia, límites y cobranza de excedentes", "El modelo de $800 + crédito se puede facturar sin intervención"],
            ["4. Multi-consultorio", "Clínicas, miembros, roles y varias agendas", "Un owner puede administrar su equipo con permisos"],
            ["5. Voz con Vapi", "Llamadas para consultar, reservar y modificar citas", "Integración auditada y con costo por uso controlado"],
        ],
        widths=[3.3, 9.3, 5.2],
    )

    add_heading(doc, "Decisiones pendientes antes de vender", 1)
    for item in [
        "Confirmar la unidad exacta de cobro de Sent.dm y convertirla en una métrica visible para el doctor.",
        "Definir el precio del excedente: tarifa del proveedor, margen de Dopilot, impuestos y redondeo.",
        "Decidir si el crédito no utilizado vence o si se acumula parcialmente; la recomendación inicial es que no se acumule.",
        "Definir límites de mensajes por hora, día y tipo de plantilla para evitar abuso y costos inesperados.",
        "Aprobar consentimiento, opt-out y políticas de privacidad para mensajes de salud.",
        "Validar con cinco doctores si $800 MXN al mes comunica claramente el valor antes de construir planes adicionales.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Métricas de negocio", 1)
    add_table(
        doc,
        ["Métrica", "Qué indica", "Uso"],
        [
            ["Activación", "Doctor que publica servicio y disponibilidad", "Mide el primer valor real"],
            ["Citas confirmadas", "Reservas que llegan a estado confirmado", "Mide utilidad de la agenda"],
            ["Uso de crédito", "Porcentaje del crédito Sent.dm consumido", "Ajusta límites y precio"],
            ["Excedente mensual", "Consumo adicional facturado", "Mide expansión de ingresos"],
            ["Retención", "Licencias activas después de 3 y 6 meses", "Valida producto y precio"],
            ["Ausencias", "Citas no atendidas antes y después", "Prueba el valor de recordatorios"],
        ],
        widths=[4.1, 8.0, 5.7],
    )

    add_heading(doc, "Conclusión", 1)
    add_body(doc, "Dopilot debe salir al mercado con una propuesta fácil de entender y de operar: $800 MXN al mes por doctor, con $150 MXN de automatizaciones Sent.dm incluidos y cobro transparente por excedente. La prioridad no es prometer una plataforma hospitalaria desde el primer día; es hacer que un doctor pueda administrar su agenda y comunicarse mejor con sus pacientes desde un solo lugar.")
    add_body(doc, "La voz con Vapi, la operación multi-consultorio y las integraciones más profundas pueden aumentar el valor después de validar el uso real. Esta secuencia protege el margen, reduce riesgo técnico y permite que cada nueva capacidad tenga un precio comprensible.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
